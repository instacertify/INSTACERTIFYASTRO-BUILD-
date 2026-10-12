"""Convert the CERTKO library workbooks into JSON for `php artisan standards:import`.

Usage: python scripts/standards-from-xlsx.py "C:/path/to/CERTKO LIBRAR"
Writes database/data/standards/{isi,crs,upcoming,bee,labs}.json (requires: pip install openpyxl).
"""

import json
import re
import sys
from collections import defaultdict
from pathlib import Path

import openpyxl

SOURCES = {
    "master": "BIS_Master_HSN_Category_QCO_Status_v5_SchemeI_and_II_Updated.xlsx",
    "bee": "BEE_Star_Label_Master_Single_Sheet_2026.xlsx",
    "labs": "BIS_Labs_Category_Scope_Wise.xlsx",
}
OUT = Path(__file__).resolve().parent.parent / "database" / "data" / "standards"


def clean(value):
    if value is None:
        return None
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    text = str(value).replace("\ufffd", "–").replace("\u00a0", " ")
    text = re.sub(r"\s+", " ", text).strip()
    return text or None


def number(value):
    text = clean(value)
    if text is None:
        return None
    try:
        n = float(text.replace(",", ""))
    except ValueError:
        return None
    return int(n) if n.is_integer() else n


def rows(book, sheet):
    return [[clean(c) for c in r] for r in book[sheet].iter_rows(values_only=True)]


def isi(book):
    out = []
    for r in rows(book, "ISI Marking Fees (Scheme-I)")[1:]:
        if not r[1]:
            continue
        slabs = [[number(r[i]), r[i + 1]] for i in (13, 15, 17) if r[i] is not None]
        out.append({
            "is": r[1], "title": r[2], "broad": r[3], "hsn4": r[4], "hsn8": r[5], "status": r[6], "order": r[7],
            "fees": {"large": number(r[8]), "medium": number(r[9]), "small": number(r[10]), "micro": number(r[11])},
            "unit": r[12], "slabs": slabs,
        })
    return out


def crs(book):
    out = []
    for r in rows(book, "CRS Marking Fees (Scheme-II)")[1:]:
        if not r[1]:
            continue
        out.append({
            "sno": number(r[0]), "product": r[1], "is": r[2], "broad": r[3], "hsn4": r[4], "hsn8": r[5], "status": r[6],
            "order": r[7], "fees": {"micro": number(r[8]), "small": number(r[9]), "large": number(r[10])}, "remarks": r[12],
        })
    return out


def upcoming(book):
    out = []
    for r in rows(book, "Upcoming QCOs (Jul-2026)")[4:]:
        if r[0] and r[2]:
            out.append({"ministry": r[1], "product": r[2], "is": r[3], "date": r[4], "scheme": r[5]})
    return out


def bee(book):
    data, section = {"mandatory": [], "voluntary": [], "fees": [], "notes": []}, None
    for r in rows(book, "BEE Master Data 2026"):
        first = r[0] or ""
        if first.startswith("A. MANDATORY"):
            section = "mandatory"
        elif first.startswith("B. VOLUNTARY"):
            section = "voluntary"
        elif first.startswith("C. GOVERNMENT"):
            section = "fees"
        elif first.startswith("NOTES"):
            data["notes"].append(first)
        elif section and first.isdigit():
            if section == "fees":
                data["fees"].append({"head": r[1], "msme": r[2], "large": r[3], "frequency": r[4], "remarks": r[5]})
            else:
                data[section].append({
                    "product": r[1], "regime": r[2], "testStandard": r[3], "ratingTable": r[4], "testingPrice": r[5],
                    "labs": r[6], "labelFee": r[7], "registrationFee": r[8],
                })
    return data


def labs(book):
    by_standard = defaultdict(list)
    for r in rows(book, "Category-Scope-wise Data")[1:]:
        if r[1] and r[3]:
            by_standard[r[1]].append({"lab": r[3], "city": r[4], "state": r[5], "price": number(r[6]), "scope": r[2]})
    return by_standard


def main():
    folder = Path(sys.argv[1]) if len(sys.argv) > 1 else Path.home() / "OneDrive/Documents/CERTKO LIBRAR"
    OUT.mkdir(parents=True, exist_ok=True)
    master = openpyxl.load_workbook(folder / SOURCES["master"], read_only=True, data_only=True)
    outputs = {
        "isi": isi(master),
        "crs": crs(master),
        "upcoming": upcoming(master),
        "bee": bee(openpyxl.load_workbook(folder / SOURCES["bee"], read_only=True, data_only=True)),
        "labs": labs(openpyxl.load_workbook(folder / SOURCES["labs"], read_only=True, data_only=True)),
    }
    for name, data in outputs.items():
        (OUT / f"{name}.json").write_text(json.dumps(data, ensure_ascii=False, indent=1), encoding="utf-8")
        print(f"{name}.json: {len(data)} records")


if __name__ == "__main__":
    main()
